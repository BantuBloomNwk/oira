#!/usr/bin/env python3
"""Pull the Oira study data from the server into one spreadsheet.

    python3 study-service/export-xlsx.py            -> ~/oira-study/Oira-Study-<date>.xlsx

Reads the export token from ~/.oira-study-token (never printed). Four sheets:
  Summary   the same totals the public live page shows, plus when this was pulled
  Women     one row per submission: shade she picked, phone, country, repeat, averages
  Readings  one row per scan
  Bench     the team's operator-tool trials, one row per trial, every field flattened
"""
import json, os, statistics, sys, urllib.request
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = 'https://app.myoira.com/api/study'
TOKEN = open(os.path.expanduser('~/.oira-study-token')).read().strip()
OUT = os.path.expanduser('~/oira-study')


def get(path, auth=True):
    req = urllib.request.Request(BASE + path, headers={'authorization': f'Bearer {TOKEN}'} if auth else {})
    with urllib.request.urlopen(req, timeout=30) as r:
        return r.read().decode()


def lines(text):
    return [json.loads(l) for l in text.splitlines() if l.strip()]


def flat(d, prefix=''):
    out = {}
    for k, v in d.items():
        key = f'{prefix}{k}'
        if isinstance(v, dict):
            out.update(flat(v, key + '.'))
        else:
            out[key] = v if not isinstance(v, list) else json.dumps(v)
    return out


def sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='43203B')
    for r in rows:
        ws.append([r.get(h) for h in header] if isinstance(r, dict) else r)
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(str(h)) + 4))
    ws.freeze_panes = 'A2'
    return ws


def main():
    count = json.loads(get('/count', auth=False))
    subs = lines(get('/export'))
    bench = lines(get('/bench/export'))

    wb = Workbook()
    wb.remove(wb.active)
    pulled = datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')
    sheet(wb, 'Summary', ['What', 'Value'], [
        ['Pulled', pulled],
        ['Women (repeats on the same phone not counted)', count.get('women', count['submissions'])],
        ['Submissions', count['submissions']],
        ['Readings', count['readings']],
        ['Phone models', count['phoneModels']],
        ['First submission', count.get('first')],
        ['Last submission', count.get('last')],
        ['By shade picked', json.dumps(count['byShade'])],
        ['By phone', json.dumps(count['byPlatform'])],
        ['By maker', json.dumps(count.get('byMake', {}))],
        ['By country (under 3 folded into other)', json.dumps(count['byCountry'])],
        ['By day', json.dumps(count['byDay'])],
        ['Team bench trials', len(bench)],
    ])

    women = []
    for s in subs:
        rs = s['readings']
        mean = lambda k: round(statistics.fmean(r[k] for r in rs), 2)
        women.append({
            'id': s['id'], 'at': s['at'], 'shadePicked': s['monkSelf'], 'platform': s['platform'],
            'make': s.get('make'), 'model': s['model'], 'modelSelf': s.get('modelSelf'), 'camera': s.get('camera'),
            'screen': s.get('screen'), 'browser': s.get('browser'), 'setting': s.get('setting'), 'localHour': s.get('localHour'),
            'country': s['country'], 'repeat': s.get('repeat', False), 'scans': len(rs),
            'meanL': mean('L'), 'meanA': mean('a'), 'meanB': mean('b'), 'meanITA': mean('ita'),
            'appShadeMedian': statistics.median(r['monk'] for r in rs),
            'meanLightK': round(statistics.fmean(r['cct'] for r in rs if r['cct'])) if any(r['cct'] for r in rs) else None,
            'meanConfidence': mean('confidence'),
        })
    sheet(wb, 'Women', list(women[0].keys()) if women else ['id'], women)

    readings = [{'submission': s['id'], 'at': s['at'], 'shadePicked': s['monkSelf'], 'model': s['model'], 'scan': i + 1, **r}
                for s in subs for i, r in enumerate(s['readings'])]
    sheet(wb, 'Readings', list(readings[0].keys()) if readings else ['submission'], readings)

    flat_bench = [flat(b) for b in bench]
    keys = []
    for b in flat_bench:
        for k in b:
            if k not in keys:
                keys.append(k)
    sheet(wb, 'Bench', keys or ['batch'], flat_bench)

    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, f"Oira-Study-{datetime.now().strftime('%Y-%m-%d-%H%M')}.xlsx")
    wb.save(path)
    print(f"{path}\n{count.get('women', count['submissions'])} women, {count['readings']} readings, {len(bench)} bench trials")


if __name__ == '__main__':
    sys.exit(main())
